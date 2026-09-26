from io import BytesIO
import json

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
import pytest
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

from app.database import get_db
from app.main import create_app
from app.models import Interview


@pytest.fixture
def client(tmp_path):
    engine = create_engine(
        f"sqlite:///{(tmp_path / 'test.db').as_posix()}",
        connect_args={"check_same_thread": False},
    )
    sessions = sessionmaker(bind=engine)
    app = create_app(engine)

    def test_db():
        with sessions() as session:
            yield session

    app.dependency_overrides[get_db] = test_db
    with TestClient(app) as test_client:
        test_client.test_sessions = sessions
        yield test_client
    engine.dispose()


def docx_bytes():
    document = Document()
    document.add_paragraph("Alex Candidate — Backend developer with Python and SQL experience.")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "FastAPI project delivery"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def create(client, filename="candidate.docx", content=None, **overrides):
    data = {
        "job_title": "Backend Developer",
        "job_description": "Build reliable Python APIs.",
        "required_skills": json.dumps(["Python", "FastAPI", "SQL"]),
    }
    data.update(overrides)
    return client.post(
        "/api/interviews", data=data,
        files={"resume": (filename, docx_bytes() if content is None else content)},
    )


def test_health(client):
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_unsupported_extension(client):
    response = create(client, filename="candidate.txt", content=b"resume text")
    assert response.status_code == 400
    assert "PDF or DOCX" in response.json()["detail"]


@pytest.mark.parametrize("skills", ['[]', 'not json', '{}', '[1]', '[" "]', 'null'])
def test_invalid_skills(client, skills):
    assert create(client, required_skills=skills).status_code == 422


@pytest.mark.parametrize("field", ["job_title", "job_description"])
def test_blank_fields(client, field):
    assert create(client, **{field: "   "}).status_code == 422


def test_successful_docx_creation(client):
    response = create(client)
    assert response.status_code == 201
    result = response.json()
    assert result["id"] > 0
    assert result["job_title"] == "Backend Developer"
    assert result["required_skills"] == ["Python", "FastAPI", "SQL"]
    assert result["resume_filename"] == "candidate.docx"
    assert result["resume_character_count"] > 20
    assert result["status"] == "created"
    assert "resume_text" not in result
    with client.test_sessions() as session:
        stored = session.get(Interview, result["id"])
        assert "FastAPI project delivery" in stored.resume_text
        assert json.loads(stored.required_skills) == result["required_skills"]


def test_get_metadata(client):
    created = create(client).json()
    response = client.get(f"/api/interviews/{created['id']}")
    assert response.status_code == 200
    result = response.json()
    assert all(result[key] == value for key, value in created.items())
    assert result["job_description"] == "Build reliable Python APIs."
    assert result["created_at"]
    assert "resume_text" not in result


def test_missing_interview(client):
    assert client.get("/api/interviews/999").status_code == 404


def test_oversized_resume(client):
    assert create(client, content=b"x" * (5 * 1024 * 1024 + 1)).status_code == 413


@pytest.mark.parametrize("filename", ["bad.pdf", "bad.docx"])
def test_corrupt_document(client, filename):
    response = create(client, filename=filename, content=b"invalid document")
    assert response.status_code == 400
    assert "Could not read" in response.json()["detail"]
    with client.test_sessions() as session:
        assert session.scalar(select(Interview)) is None


def test_pdf_without_text(client):
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    stream = BytesIO()
    writer.write(stream)
    response = create(client, filename="scanned.pdf", content=stream.getvalue())
    assert response.status_code == 400
    assert "Could not extract readable text from this PDF" in response.json()["detail"]


def test_pdf_with_text(client):
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})
    })
    contents = DecodedStreamObject()
    contents.set_data(b"BT /F1 12 Tf 50 700 Td (Alex Candidate Python SQL backend developer) Tj ET")
    page[NameObject("/Contents")] = contents
    stream = BytesIO()
    writer.write(stream)
    response = create(client, filename="candidate.PDF", content=stream.getvalue())
    assert response.status_code == 201
    assert response.json()["resume_character_count"] > 20


@pytest.mark.parametrize("origin", [
    "http://localhost:5173", "http://127.0.0.1:5173",
    "http://127.0.0.1:5500", "http://localhost:5501",
])
def test_cors(client, origin):
    response = client.options("/api/interviews", headers={
        "Origin": origin, "Access-Control-Request-Method": "POST",
    })
    assert response.headers["access-control-allow-origin"] == origin


def test_cors_rejects_external_origin(client):
    response = client.options("/api/interviews", headers={
        "Origin": "http://example.com:5500", "Access-Control-Request-Method": "POST",
    })
    assert response.status_code == 400
    assert "access-control-allow-origin" not in response.headers
