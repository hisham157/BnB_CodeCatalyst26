from io import BytesIO
from pathlib import Path
import re

from docx import Document
from pypdf import PdfReader


class DocumentExtractionError(ValueError):
    """A resume cannot be read or has insufficient readable text."""


def extract_resume_text(filename: str, file_bytes: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in {".pdf", ".docx"}:
        raise DocumentExtractionError("Unsupported resume type. Upload a PDF or DOCX file.")

    try:
        stream = BytesIO(file_bytes)
        if extension == ".pdf":
            reader = PdfReader(stream)
            parts = [page.extract_text() or "" for page in reader.pages]
        else:
            document = Document(stream)
            parts = [paragraph.text for paragraph in document.paragraphs]
            parts.extend(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            )
    except Exception as exc:
        raise DocumentExtractionError(
            "Could not read this resume. Upload a valid, unencrypted PDF or DOCX file."
        ) from exc

    text = "\n".join(
        line for part in parts for raw in part.splitlines()
        if (line := re.sub(r"\s+", " ", raw).strip())
    )
    if sum(character.isalnum() for character in text) < 20:
        label = "PDF" if extension == ".pdf" else "DOCX"
        raise DocumentExtractionError(
            f"Could not extract readable text from this {label}. "
            "Use a text-based resume; scanned documents require OCR, which is not supported."
        )
    return text
